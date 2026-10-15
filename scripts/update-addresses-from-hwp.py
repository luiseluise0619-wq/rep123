"""Reconcile the existing map with the final HWP classification, without geocoding.

Inputs are local review artifacts, not uploaded to the public repository. Raw
address strings and coordinates remain unchanged. Only the classification tree
and business-count metadata are updated. Requires openpyxl for read-only input.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SHORT = {
    '서울특별시': '서울', '부산광역시': '부산', '대구광역시': '대구',
    '인천광역시': '인천', '광주광역시': '광주', '대전광역시': '대전',
    '울산광역시': '울산', '세종특별자치시': '세종', '경기도': '경기',
    '강원특별자치도': '강원', '충청북도': '충북', '충청남도': '충남',
    '전북특별자치도': '전북', '전라남도': '전남', '경상북도': '경북',
    '경상남도': '경남', '제주특별자치도': '제주',
}
FULL = {v: k for k, v in SHORT.items()}
MAIN = {'서울', '경기', '인천', '부산', '대구', '대전', '청주'}


def canonical(address):
    address = address.strip()
    if address.startswith('전남광주통합특별시 '):
        district = address.split()[1]
        province = '광주광역시' if district in {'광산구', '남구', '동구', '북구', '서구'} else '전라남도'
        return province + ' ' + address.split(' ', 1)[1]
    return address


def target_region(address):
    if address.startswith('충청북도 청주시 '):
        return '청주'
    label = SHORT.get(address.split()[0])
    return label if label in MAIN | {'전북', '전남', '경북', '경남'} else None


def read_web(path):
    text = path.read_text(encoding='utf-8')
    prefix = 'window.ADDR_DATA = '
    assert text.startswith(prefix), 'Unexpected addresses.js format'
    tree = json.loads(text[len(prefix):].strip().removesuffix(';'))
    return [(s, g, d, p) for s, gs in tree.items() for g, ds in gs.items()
            for d, ps in ds.items() for p in ps]


def cells(table):
    return [[''.join(c.itertext()).strip() for c in row.findall('./CELL/PARALIST')]
            for row in table.findall('./ROW')]


def number(text):
    return int(re.sub(r'[\s,]', '', text))


def norm(text):
    return re.sub(r'\s+', '', text)


def verify_hwp(xml_path, audit):
    # Hancom's exported text is UTF-8 with a BOM but declares UTF-16.
    root = ET.fromstring(xml_path.read_text(encoding='utf-8-sig'))
    national = Counter()
    first_table = root.find('.//TABLE')
    for row in cells(first_table)[1:]:
        for offset in (0, 2, 4):
            label = norm(row[offset])
            if label != '합계':
                national[label] += number(row[offset + 1])
            else:
                assert number(row[offset + 1]) == 44075
    groups, districts, unknown = Counter(), Counter(), Counter()
    current = None
    for para in root.findall('./BODY/SECTION/P'):
        title = ''.join(c.text or '' for c in para.findall('./TEXT/CHAR'))
        match = re.match(r'^(서울|경기|인천|부산|대구|대전|청주) 행정동별 사업자 수', title)
        if match:
            current = match.group(1)
        if title.startswith('행정동 미확정 3개'):
            current = None
        if not current:
            continue
        for table in para.findall('.//TABLE'):
            rows = cells(table)
            if len(rows[0]) != 6 or norm(rows[0][0]) != '시군구':
                continue
            for row in rows[1:]:
                for offset in (0, 3):
                    district, dong, value = row[offset:offset + 3]
                    if not district:
                        continue
                    if district.endswith(' 합계'):
                        districts[(current, norm(district.removesuffix(' 합계')))] += number(value)
                    elif dong == '행정동 미확정':
                        unknown[(current, norm(district))] += number(value)
                    else:
                        groups[(current, norm(district), norm(dong))] += number(value)
    expected = Counter({(g['region'], norm(g['district']), norm(g['dong'])): g['businesses']
                        for g in audit['groups'] if g['region'] in MAIN})
    expected_districts = Counter({(d['region'], norm(d['district'])): d['businesses']
                                  for d in audit['district_totals'] if d['region'] in MAIN})
    expected_unknown = Counter((r['region'], norm(r['district'])) for r in audit['unresolved'])
    assert groups == expected, 'HWP administrative-dong table differs from review records'
    assert districts == expected_districts, 'HWP district subtotals differ'
    assert unknown == expected_unknown, 'HWP unresolved counts differ'
    assert len(groups) == 1661 and sum(districts.values()) == 27914
    return national


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--audit', type=Path, required=True)
    parser.add_argument('--hwp-xml', type=Path, required=True)
    parser.add_argument('--hwp', type=Path, required=True)
    parser.add_argument('--source', type=Path, default=ROOT / '원본_260810.xlsx')
    args = parser.parse_args()
    audit = json.loads(args.audit.read_text(encoding='utf-8'))
    hwp_bytes = args.hwp.read_bytes()
    hwp_hash = hashlib.sha256(hwp_bytes).hexdigest()
    # Fingerprint of the native HWP whose saved HWPML and PDF were verified.
    assert hwp_hash == '1defab1677b52671bc8e7f8153c2200a99936ed148c1f9db6d6c4326e3b953e1', 'HWP changed: re-export and verify it first'
    national = verify_hwp(args.hwp_xml, audit)

    from openpyxl import load_workbook
    workbook = load_workbook(args.source, read_only=True, data_only=True)
    source = Counter()
    source_national = Counter()
    source_rows = 0
    for row in workbook['Sheet2'].iter_rows(min_row=4, max_col=6, values_only=True):
        if not row[4]:
            continue
        address = canonical(str(row[4]))
        count = int(row[5])
        assert count > 0 and count == row[5]
        source[address] += count
        source_national[SHORT[address.split()[0]]] += count
        source_rows += 1
    workbook.close()
    assert source_national == national and sum(source.values()) == 44075

    old = read_web(ROOT / 'web/addresses.js')
    old_counts = Counter()
    for _, _, _, p in old:
        old_counts[canonical(p[2])] += p[3]
    assert old_counts == source and len(old) == source_rows == 42401
    by_address = defaultdict(list)
    for record in audit['business_records']:
        by_address[record['address']].append(record)
    target_counts = Counter({a: len(rs) for a, rs in by_address.items()})
    assert target_counts == Counter({a: n for a, n in source.items() if target_region(a)})

    tree = {}
    web_groups, web_districts, web_regions = Counter(), Counter(), Counter()
    address_counts = Counter()
    changed_addresses = changed_businesses = 0
    for old_s, old_g, old_d, p in old:
        address = canonical(p[2])
        s = address.split()[0]
        g, d, code, scope = old_g, old_d, '', ''
        if address in by_address:
            records = by_address[address]
            assignments = {(r['district'], r['dong'], r['code'], r['region']) for r in records}
            assert len(assignments) == 1, 'One source address has conflicting administrative assignments'
            g, d, code, scope = next(iter(assignments))
            assert p[3] == len(records)
            web_groups[(scope, g, d, code)] += p[3]
            web_districts[(scope, g)] += p[3]
            web_regions[scope] += p[3]
            address_counts[scope] += 1
        if (old_s, old_g, old_d) != (s, g, d):
            changed_addresses += 1
            changed_businesses += p[3]
        # The original address and coordinate quality are intentionally preserved.
        new_p = p[:5] + [code, scope]
        tree.setdefault(s, {}).setdefault(g, {}).setdefault(d, []).append(new_p)

    expected_groups = Counter({(g['region'], g['district'], g['dong'], g['code']): g['businesses']
                               for g in audit['groups']})
    expected_groups.update((r['region'], r['district'], r['dong'], '') for r in audit['unresolved'])
    assert web_groups == expected_groups
    assert web_districts == Counter({(g['region'], g['district']): g['businesses'] for g in audit['district_totals']})
    assert web_regions == Counter({c['region']: c['source_businesses'] for c in audit['controls']})
    assert sum(web_regions.values()) == 36452 and len(web_districts) == 199
    flat = [(s, g, d, p) for s, gs in tree.items() for g, ds in gs.items() for d, ps in ds.items() for p in ps]
    after = Counter()
    for s, _, _, p in flat:
        after[SHORT[s]] += p[3]
    assert after == national
    assert Counter((p[2], p[3]) for _, _, _, p in flat) == Counter((p[2], p[3]) for _, _, _, p in old)
    assert Counter((p[2], *p[:2], p[4]) for _, _, _, p in flat) == Counter((p[2], *p[:2], p[4]) for _, _, _, p in old)
    assert len(tree) == 17 and len(flat) == 42401

    controls = []
    for c in audit['controls']:
        scope = c['region']
        controls.append({
            'region': scope, 'province': FULL.get(scope, '충청북도'),
            'businesses': c['source_businesses'], 'addresses': address_counts[scope],
            'assigned': c['assigned'], 'unresolved': c['unresolved'],
            'districts': sum(k[0] == scope for k in web_districts),
            'dongs': sum(g['region'] == scope for g in audit['groups']),
        })
    yp = [r for r in audit['business_records'] if r['address'].startswith('서울특별시 영등포구 ')]
    yp_same = sum(r['district'] == '영등포구' for r in yp)
    yp_current = web_districts[('서울', '영등포구')]
    assert len(yp) == 534 and yp_same == 533
    meta = {
        'asOf': '2026-10-15', 'sourceDocument': args.hwp.name, 'sourceSha256': hwp_hash,
        'nationalBusinesses': 44075, 'nationalAddresses': len(flat),
        'targetBusinesses': 36452, 'targetAddresses': sum(address_counts.values()),
        'assignedBusinesses': 36449, 'unresolvedBusinesses': 3,
        'targetDistricts': 199, 'mainDongGroups': 1661,
        'controls': controls,
        'unresolved': [{'address': r['address'], 'province': FULL[r['region']],
                        'district': r['district'], 'businesses': r['businesses']} for r in audit['unresolved']],
        'originalYeongdeungpo': {'businesses': len(yp), 'sameOriginalDistrictBusinesses': yp_same,
                                  'currentDistrictBusinesses': yp_current,
                                  'incomingFromOtherOriginalDistricts': yp_current - yp_same,
                                  'otherDistrictBusinesses': len(yp) - yp_same,
                                  'otherDistrict': '동작구 신대방1동'},
        'userProvidedClassifications': 22,
        'notes': [
            '사업자수는 원자료 주소별 사업자수의 합계이며 주소 건수와 다릅니다.',
            '대상 11개 지역은 최종 한글 집계와 대조했습니다. 그 외 지역은 기존 주소 분류를 유지합니다.',
            '원주소는 보존합니다. 목록의 시군구·행정동은 한글 집계 관할이라 원주소 표기와 다를 수 있습니다.',
            '사용자가 제공한 22개 행정동 분류가 포함되며, 독립 검증 완료를 의미하지 않습니다.',
            '지도 좌표는 기존 값을 유지했습니다. 행정동 분류 확정과 좌표 정확성은 별개입니다.',
        ],
        'checks': {'hwpMainDongTables': True, 'hwpDistrictSubtotals': True,
                   'sourceAddressesPreserved': True, 'coordinatesPreserved': True,
                   'sourceBusinessCounts': True, 'missingBusinesses': 0, 'excessBusinesses': 0},
    }
    (ROOT / 'web/addresses.js').write_text('window.ADDR_DATA = ' + json.dumps(tree, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
    (ROOT / 'web/data-meta.js').write_text('window.ADDR_META = ' + json.dumps(meta, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
    (ROOT / 'web/data-validation.json').write_text(json.dumps(meta, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    summary = {k: meta[k] for k in ['nationalBusinesses', 'nationalAddresses', 'targetBusinesses', 'targetAddresses', 'targetDistricts', 'mainDongGroups', 'unresolvedBusinesses', 'checks']}
    summary.update(updatedAddresses=changed_addresses, updatedBusinesses=changed_businesses)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
