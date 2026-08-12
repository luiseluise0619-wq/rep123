# -*- coding: utf-8 -*-
"""
전화번호·사업자등록번호를 '좌르륵 나열'만 하면 주소 순서에 맞춰 자동 입력
================================================================================
주소는 그대로 두고, 전화번호 목록과 사업자등록번호 목록을 '행 순서(위→아래)'대로
각 주소에 자동으로 짝지어 넣어줍니다.

  주소[0] ↔ 전화[0] ↔ 사업자[0]
  주소[1] ↔ 전화[1] ↔ 사업자[1]
  ...

결과 파일(주소_연락처_병합.xlsx)은 입력 규약(A=주소, D=전화, E=사업자)에 맞춰 저장되므로,
이어서 `주소_분류_및_지도생성.py` 를 실행하면 G열(전화)·H열(사업자)까지 자동 정리됩니다.

■ 전화/사업자 목록을 넣는 방법 (아무거나 택1)
  (1) 텍스트 파일 : 한 줄에 하나씩 (전화번호.txt / 사업자등록번호.txt)
  (2) 엑셀 파일   : 첫 번째 열에 위→아래로 나열
  (3) 이 파일 안에 직접 붙여넣기 : 아래 TEL_PASTE / BIZ_PASTE 삼중따옴표 안에
"""

import os
from openpyxl import load_workbook, Workbook

# =============================================================================
# CONFIG (경로는 윈도우 역슬래시 방지 r'')
# =============================================================================
BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
ADDR_XLSX  = os.path.join(BASE_DIR, r'exceldata1.xlsx')   # A열에 주소
OUT_XLSX   = os.path.join(BASE_DIR, r'주소_연락처_병합.xlsx')
ADDR_COL   = 1            # 주소가 있는 열(A=1)
ADDR_HEADER = False       # 주소 파일 첫 행이 머리글이면 True

# 전화/사업자 소스 파일 (없으면 None). .txt(한 줄당 하나) 또는 .xlsx(첫 열)
TEL_SOURCE = os.path.join(BASE_DIR, r'전화번호.txt')
BIZ_SOURCE = os.path.join(BASE_DIR, r'사업자등록번호.txt')

# 또는 여기에 직접 붙여넣기(줄바꿈으로 구분). 파일보다 우선 적용됨.
TEL_PASTE = r"""
"""
BIZ_PASTE = r"""
"""

# 길이가 주소 수보다 짧으면 빈칸으로 채우고, 길면 잘라냄(경고 출력)
# =============================================================================


def read_addresses(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if ADDR_HEADER and i == 0:
            continue
        v = row[ADDR_COL - 1] if len(row) >= ADDR_COL else None
        if v is None or str(v).strip() == '':
            continue
        out.append(str(v).strip())
    wb.close()
    return out


def read_list(paste, source):
    """붙여넣기(우선) 또는 파일(.txt/.xlsx)에서 값 목록을 읽는다."""
    # 1) 직접 붙여넣기 우선
    vals = [ln.strip() for ln in (paste or '').splitlines() if ln.strip() != '']
    if vals:
        return vals
    # 2) 파일
    if source and os.path.exists(source):
        if source.lower().endswith(('.xlsx', '.xlsm')):
            wb = load_workbook(source, read_only=True, data_only=True)
            ws = wb.active
            out = []
            for row in ws.iter_rows(values_only=True):
                v = row[0] if row else None
                if v is not None and str(v).strip() != '':
                    out.append(str(v).strip())
            wb.close()
            return out
        with open(source, 'r', encoding='utf-8') as f:
            return [ln.strip() for ln in f if ln.strip() != '']
    return []


def attach_contacts(addresses, tels, bizs):
    """
    핵심 함수: 주소 순서에 맞춰 전화·사업자를 정렬해 (주소, 전화, 사업자) 목록 반환.
    부족분은 '' 로 채우고, 초과분은 무시한다.
    """
    n = len(addresses)

    def fit(lst, label):
        if len(lst) > n:
            print('  ⚠ %s %d개가 주소(%d)보다 많아 뒤쪽을 잘라냅니다.' % (label, len(lst), n))
        elif len(lst) < n:
            print('  ⚠ %s %d개가 주소(%d)보다 적어 나머지는 빈칸으로 둡니다.' % (label, len(lst), n))
        return (lst + [''] * n)[:n]

    tels = fit(tels, '전화번호')
    bizs = fit(bizs, '사업자등록번호')
    return [(addresses[i], tels[i], bizs[i]) for i in range(n)]


def save_merged(rows, path):
    """A=주소, D=전화, E=사업자 레이아웃으로 저장(입력 규약과 동일)."""
    wb = Workbook()
    ws = wb.active
    ws.title = '주소_연락처'
    ws.append(['주소', '', '', '전화번호', '사업자등록번호'])  # A,(B,C 비움),D,E
    for addr, tel, biz in rows:
        ws.append([addr, None, None, tel, biz])
    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 16
    wb.save(path)


def main():
    print('[1] 주소 읽기 :', ADDR_XLSX)
    addresses = read_addresses(ADDR_XLSX)
    print('    → 주소 %d건' % len(addresses))

    tels = read_list(TEL_PASTE, TEL_SOURCE)
    bizs = read_list(BIZ_PASTE, BIZ_SOURCE)
    print('[2] 전화 %d개 · 사업자 %d개 읽음' % (len(tels), len(bizs)))

    print('[3] 주소 순서에 맞춰 자동 정렬 중...')
    rows = attach_contacts(addresses, tels, bizs)

    save_merged(rows, OUT_XLSX)
    filled_tel = sum(1 for _, t, _ in rows if t)
    filled_biz = sum(1 for _, _, b in rows if b)
    print('[완료] 저장:', OUT_XLSX)
    print('    → 전화 채움 %d건 / 사업자 채움 %d건' % (filled_tel, filled_biz))
    print('    → 이어서  python 주소_분류_및_지도생성.py  실행 시 G·H열까지 정리됩니다.')


if __name__ == '__main__':
    main()
