# -*- coding: utf-8 -*-
"""
피벗(요약) 형태의 원본 엑셀 → 표준 레이아웃(exceldata1.xlsx)으로 정규화
================================================================================
원본 형식(피벗을 '전체 펼침'한 상태):
    · 0~2행: 머리말/헤더  (대분류·중분류·담당센터·시군구·주소·사업장 수)
    · E열(5): 주소,  F열(6): 사업장 수
    · 주소가 빈 행(소계/총합계)은 건너뜀

출력 exceldata1.xlsx:
    · A열: 주소,  C열: 사업장 수  (D=전화, E=사업자용은 비워둠)
※ 피벗이 접혀 있으면(구 앞 +표시) 개별 주소가 파일에 없으니, 반드시 '전체 펼침' 후 저장하세요.
"""
import os
from openpyxl import load_workbook, Workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(BASE_DIR, r'원본_260810.xlsx')     # 원본(피벗 펼친 파일)
OUT  = os.path.join(BASE_DIR, r'exceldata1.xlsx')       # 표준 레이아웃 출력
SKIP_ROWS = 3          # 머리말/헤더 행 수
COL_ADDR_SRC = 5       # 원본 주소 열(E)
COL_CNT_SRC  = 6       # 원본 사업장수 열(F)


def main():
    wb = load_workbook(SRC, data_only=True)
    ws = wb.active
    out = Workbook(); o = out.active; o.title = '주소'
    n = 0; biz = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < SKIP_ROWS:
            continue
        addr = row[COL_ADDR_SRC - 1] if len(row) >= COL_ADDR_SRC else None
        if addr is None or str(addr).strip() == '':
            continue                       # 소계/총합계 등 주소 없는 행 제외
        cnt_raw = row[COL_CNT_SRC - 1] if len(row) >= COL_CNT_SRC else None
        try:
            cnt = int(float(cnt_raw)) if str(cnt_raw).strip() != '' else 1
        except (ValueError, TypeError):
            cnt = 1
        o.append([str(addr).strip(), None, cnt])   # A=주소, B빈칸, C=사업장수
        n += 1; biz += cnt
    out.save(OUT)
    print('정규화 완료: 주소 %d건, 총 사업장 %d개 → %s' % (n, biz, OUT))


if __name__ == '__main__':
    main()
